"""BMC → BMC Helix address reconciliation (rapprochement_bmc_helix_historique.xlsx).

Many BMC people moved from @bmc.com to a BMC Helix address. The reconciliation workbook lists, per original contact,
what was found in mail history. Its sheets (French names; matched without accents or case):

  Correspondances   probable match on the full name: original @bmc.com address → observed Helix address
  Non trouvés       no Helix address found: keep the original address (nothing is deduced)
  Ambigus           several Helix addresses for one name: never applied, only reported
  Déjà Helix        the original file already had a Helix address

Applying it, per row:
  match         the contact with the old address takes the new one (the old one is kept in extra.previous_emails);
                if a contact already has the new address too, the two are merged into it;
                if neither exists, a contact is created with the new address
  already Helix / not found   the contact is created if no contact has that address; existing ones only get blanks filled
  ambiguous     skipped

Every contact touched keeps the workbook's evidence in extra.bmc_helix, and address changes and merges are noted on
its timeline. Running the same workbook again changes nothing more.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Account, Activity, Contact, User
from ..services import find_contact_by_email, merge_contacts

SOURCE = "BMC Helix reconciliation"
# account names in the workbook that mean BMC itself
BMC_NAMES = {"bmc", "bmc software", "bmc helix", "bmc software (bmc helix)"}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


SHEETS = {"correspondances": "match", "non trouves": "not_found", "ambigus": "ambiguous", "deja helix": "already_helix"}
COLUMNS = {  # normalised header → field
    "id original": "source_id", "prenom original": "first_name", "nom original": "last_name",
    "e mail original": "email", "e mail helix observe": "helix_email", "compte original": "account",
    "statut": "status", "preuve": "evidence", "date source utc": "observed_at", "nom dans le mail": "name_in_mail",
}


@dataclass
class Row:
    kind: str  # match, not_found, ambiguous, already_helix
    line: int
    source_id: str = ""
    first_name: str = ""
    last_name: str = ""
    email: str = ""
    helix_email: str = ""
    account: str = ""
    status: str = ""
    evidence: str = ""
    observed_at: str = ""
    name_in_mail: str = ""

    @property
    def name(self) -> str:
        return " ".join(p for p in (self.first_name, self.last_name) if p) or self.email


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return str(v).strip()


def parse(raw: bytes) -> tuple[list[Row], list[str]]:
    """Rows of every recognised sheet, and warnings. Raises ValueError if this is not a reconciliation workbook."""
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read the Excel file: {exc}") from None
    rows, warnings, found = [], [], set()
    for ws in wb.worksheets:
        kind = next((k for prefix, k in SHEETS.items() if _norm(ws.title).startswith(prefix)), None)
        if not kind:
            continue
        found.add(kind)
        header, cols = None, {}
        for i, values in enumerate(ws.iter_rows(values_only=True), start=1):
            if header is None:
                names = [_norm(v) for v in values]
                if "e mail original" in names:
                    header = i
                    cols = {COLUMNS[n]: j for j, n in enumerate(names) if n in COLUMNS}
                continue
            r = Row(kind=kind, line=i, **{f: _cell(values[j]) if j < len(values) else "" for f, j in cols.items()})
            r.email, r.helix_email = r.email.lower(), r.helix_email.lower()
            if not r.email:
                continue  # notes such as "Aucune ligne dans cette catégorie."
            if "@" not in r.email or (kind == "match" and "@" not in r.helix_email):
                warnings.append(f"{ws.title} line {i}: no usable email address, skipped")
                continue
            rows.append(r)
        if header is None:
            warnings.append(f"Sheet '{ws.title}' has no 'E-mail original' column: skipped")
    if "match" not in found and "already_helix" not in found:
        raise ValueError("This does not look like a BMC Helix reconciliation workbook (no 'Correspondances' or "
                         "'Déjà Helix' sheet)")
    return rows, warnings


@dataclass
class Report:
    dry_run: bool
    counts: Counter = field(default_factory=Counter)
    rows_by_sheet: Counter = field(default_factory=Counter)
    changes: list[dict] = field(default_factory=list)  # one line per contact changed, for the preview
    warnings: list[str] = field(default_factory=list)

    def add(self, action: str, r: Row, contact_id: int | None = None, detail: str = ""):
        self.counts[action] += 1
        if action != "unchanged":
            self.changes.append({"action": action, "sheet": r.kind, "name": r.name, "email": r.email,
                                 "new_email": r.helix_email or None, "account": r.account or None,
                                 "contact_id": contact_id, "detail": detail})

    def as_dict(self) -> dict:
        return {"dry_run": self.dry_run, "rows_by_sheet": dict(self.rows_by_sheet), "counts": dict(self.counts),
                "changes": self.changes[:2000], "warnings": self.warnings[:200]}


class _Accounts:
    """Workbook account name → CRM account. BMC names go to the BMC platform account; others only if they exist."""

    def __init__(self, db: Session):
        self.db, self.cache = db, {}
        self.bmc = db.scalar(select(Account).where(Account.website.ilike("%bmc.com%"))
                             .order_by(Account.account_type != "platform_partner", Account.id)) \
            or db.scalar(select(Account).where(func.lower(Account.name).in_(BMC_NAMES)).order_by(Account.id))

    def get(self, name: str) -> Account | None:
        key = name.strip().lower()
        if not key:
            return None
        if key in BMC_NAMES:
            return self.bmc
        if key not in self.cache:
            self.cache[key] = self.db.scalar(select(Account).where(func.lower(Account.name) == key)
                                             .order_by(Account.id))
        return self.cache[key]


def _evidence(r: Row) -> dict:
    return {k: v for k, v in {"sheet": r.kind, "status": r.status, "evidence": r.evidence,
                              "observed_at": r.observed_at, "source_id": r.source_id,
                              "original_email": r.email}.items() if v}


def _note(db: Session, c: Contact, user: User | None, subject: str, body: str):
    db.add(Activity(type="note", subject=subject, body=body, contact_id=c.id, account_id=c.account_id,
                    created_by_id=user.id if user else None, completed_at=dt.datetime.now(dt.UTC)))


def _fill(c: Contact, r: Row, accounts: _Accounts) -> list[str]:
    """Fill blank fields from the workbook; never overwrite what the CRM has."""
    filled = []
    if not c.first_name and r.first_name:
        c.first_name = r.first_name
        filled.append("first name")
    if not c.last_name and r.last_name:
        c.last_name = r.last_name
        filled.append("last name")
    if c.account_id is None and (acc := accounts.get(r.account)):
        c.account_id = acc.id
        filled.append(f"account {acc.name}")
    return filled


def _remember(c: Contact, r: Row, old_email: str | None = None):
    extra = dict(c.extra or {})
    extra["bmc_helix"] = _evidence(r)
    if old_email and old_email != (c.email or "").lower():
        prev = list(extra.get("previous_emails") or [])
        if old_email not in prev:
            prev.append(old_email)
        extra["previous_emails"] = prev
    if r.account and not c.account_id:
        extra.setdefault("company", r.account)
    c.extra = extra


def _create(db: Session, r: Row, email: str, accounts: _Accounts, owner: User | None) -> Contact:
    acc = accounts.get(r.account)
    c = Contact(first_name=r.first_name or None, last_name=r.last_name or None, email=email,
                account_id=acc.id if acc else None, source="import", lifecycle="customer" if acc else "lead",
                role="platform_account_manager" if acc is not None and acc is accounts.bmc else "customer_contact",
                owner_id=owner.id if owner else None, extra={})
    _remember(c, r, r.email if email != r.email else None)
    db.add(c)
    db.flush()
    return c


def apply(db: Session, rows: list[Row], *, dry_run: bool, user: User | None = None,
          apply_matches: bool = True, create_missing: bool = True) -> Report:
    """Apply the workbook. With dry_run the work is done in a savepoint and rolled back, so the report is exact."""
    rep = Report(dry_run=dry_run)
    accounts = _Accounts(db)
    sp = db.begin_nested()
    for r in rows:
        rep.rows_by_sheet[r.kind] += 1
        if r.kind == "ambiguous":
            rep.add("skipped_ambiguous", r, detail="several possible Helix addresses: review by hand")
            continue
        old = find_contact_by_email(db, r.email)
        if r.kind == "match":
            new = find_contact_by_email(db, r.helix_email)
            if not apply_matches:
                rep.add("skipped_match", r, (new or old).id if (new or old) else None, "probable matches not applied")
                continue
            if old and new and old.id != new.id:
                merge_contacts(db, keep=new, drop=old)
                _fill(new, r, accounts)
                _remember(new, r, r.email)
                _note(db, new, user, f"Merged with {r.email} ({SOURCE})",
                      f"{r.email} and {r.helix_email} were the same person: merged into this contact. "
                      f"{r.evidence} (seen {r.observed_at}).")
                rep.add("merged", r, new.id, f"contact {old.id} merged into {new.id}")
            elif old and not new:
                old.email = r.helix_email
                _fill(old, r, accounts)
                _remember(old, r, r.email)
                _note(db, old, user, f"Email changed to {r.helix_email} ({SOURCE})",
                      f"Was {r.email}. {r.evidence} (seen {r.observed_at}).")
                rep.add("email_updated", r, old.id, f"{r.email} → {r.helix_email}")
            elif new:
                known = r.email in ((new.extra or {}).get("previous_emails") or [])
                filled = _fill(new, r, accounts)
                _remember(new, r, r.email)
                rep.add("updated" if filled or not known else "unchanged", r, new.id,
                        ", ".join(filled) or f"already on the Helix address; {r.email} kept as a previous address")
            elif create_missing:
                c = _create(db, r, r.helix_email, accounts, user)
                rep.add("created", r, c.id, f"with the Helix address (was {r.email})")
            else:
                rep.add("skipped_missing", r)
            continue
        # already_helix / not_found: the address in the file is the one to keep
        if old:
            filled = _fill(old, r, accounts)
            _remember(old, r)  # the evidence is recorded quietly; only filled fields count as a change
            rep.add("updated" if filled else "unchanged", r, old.id, ", ".join(filled))
        elif create_missing:
            c = _create(db, r, r.email, accounts, user)
            rep.add("created", r, c.id, "no Helix address found: kept the BMC address" if r.kind == "not_found"
                    else "")
        else:
            rep.add("skipped_missing", r)
    db.flush()
    if dry_run:
        sp.rollback()
    else:
        sp.commit()
    if accounts.bmc is None:
        rep.warnings.append("No BMC account found in the CRM (website bmc.com or named BMC): BMC people were "
                            "created without an account")
    return rep
