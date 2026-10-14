"""Excel / CSV import (with column mapping and de-duplication) and export."""
import csv
import io
import json
import re
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..db import get_db
from ..enums import AccountType, ContactRole, LeadSource, Lifecycle
from ..models import Account, Brand, Contact, ImportJob, Opportunity, User
from ..schemas import ImportPreview, ImportResult
from ..security import get_current_user, require_admin, require_writer
from ..services import email_domain, find_contact_by_email, find_or_create_account
from . import accounts as accounts_router
from . import contacts as contacts_router
from . import opportunities as opps_router

router = APIRouter(prefix="/api", tags=["import-export"])
MAX_ROWS = 50_000
MAX_BYTES = 20 * 1024 * 1024

CONTACT_FIELDS = {
    "first_name": ["first name", "firstname", "given name", "prenom", "prénom"],
    "last_name": ["last name", "lastname", "surname", "family name", "nom"],
    "full_name": ["name", "full name", "contact name", "contact"],
    "email": ["email", "e-mail", "email address", "mail", "courriel"],
    "phone": ["phone", "telephone", "téléphone", "work phone", "business phone", "tel"],
    "mobile": ["mobile", "cell", "mobile phone", "portable"],
    "job_title": ["title", "job title", "position", "role", "fonction"],
    "company": ["company", "account", "organization", "organisation", "company name", "société", "societe",
                "account name"],
    "website": ["website", "domain", "company website", "url", "web"],
    "country": ["country", "pays", "country/region"],
    "linkedin": ["linkedin", "linkedin url", "linkedin profile"],
    "lifecycle": ["lifecycle", "lifecycle stage", "status", "lead status"],
    "source": ["source", "lead source"],
    "contact_role": ["contact role", "contact type", "type"],
    "brand": ["brand", "product line", "interest"],
    "consent_marketing": ["consent", "opt in", "opt-in", "marketing consent", "subscribed"],
    "notes": ["notes", "note", "comments", "comment", "description"],
}
ACCOUNT_FIELDS = {
    "name": ["name", "account", "account name", "company", "company name", "organization", "société"],
    "website": ["website", "domain", "url", "web"],
    "account_type": ["type", "account type", "category"],
    "industry": ["industry", "sector", "secteur"],
    "country": ["country", "pays"],
    "region": ["region", "state", "territory"],
    "city": ["city", "ville"],
    "address": ["address", "street", "adresse"],
    "phone": ["phone", "telephone", "tel"],
    "vat_number": ["vat", "vat number", "tax id"],
    "notes": ["notes", "description", "comments"],
}
FIELDS = {"contacts": CONTACT_FIELDS, "accounts": ACCOUNT_FIELDS}
_email = TypeAdapter(EmailStr)


def _norm(h: str) -> str:
    return re.sub(r"[\s_\-]+", " ", str(h or "").strip().lower())


def read_table(upload: UploadFile) -> tuple[list[str], list[dict]]:
    raw = upload.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "File too large (max 20 MB)")
    name = (upload.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(422, f"Could not read Excel file: {exc}") from None
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows:
            return [], []
        headers = [str(h).strip() if h is not None else f"Column {i + 1}" for i, h in enumerate(rows[0])]
        data = [{headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r[:len(headers)])}
                for r in rows[1:]]
    elif name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = [r for r in reader if any(c.strip() for c in r)]
        if not rows:
            return [], []
        headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
        data = [{headers[i]: v.strip() for i, v in enumerate(r[:len(headers)])} for r in rows[1:]]
    else:
        raise HTTPException(422, "Unsupported file type: upload .xlsx or .csv")
    if len(data) > MAX_ROWS:
        raise HTTPException(422, f"Too many rows (max {MAX_ROWS})")
    return headers, data


def suggest_mapping(entity: str, headers: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {}
    used = set()
    for field, synonyms in FIELDS[entity].items():
        options = {_norm(field.replace("_", " ")), *(_norm(s) for s in synonyms)}
        match = next((h for h in headers if _norm(h) in options and h not in used), None)
        if match:
            used.add(match)
        mapping[field] = match
    return mapping


@router.post("/imports/preview", response_model=ImportPreview)
def preview(entity: Literal["contacts", "accounts"] = Form(...), file: UploadFile = File(...),
            _: User = Depends(require_writer)):
    headers, rows = read_table(file)
    return ImportPreview(headers=headers, sample=rows[:10], total_rows=len(rows),
                         suggested_mapping=suggest_mapping(entity, headers), fields=list(FIELDS[entity]))


def _enum(value: str, enum_cls, default=None):
    v = _norm(value).replace(" ", "_")
    for member in enum_cls:
        if v in (member.value, member.name) or _norm(member.value.replace("_", " ")) == _norm(value):
            return member.value
    return default


def _truthy(v: str) -> bool:
    return _norm(v) in {"1", "true", "yes", "y", "oui", "x", "opt in", "opted in", "subscribed"}


def _domain_from(site: str | None) -> str | None:
    if not site:
        return None
    site = site.strip().lower()
    if "@" in site:
        return email_domain(site)
    site = re.sub(r"^https?://", "", site).removeprefix("www.")
    return site.split("/")[0] or None


@router.post("/imports/commit", response_model=ImportResult)
def commit(
    entity: Literal["contacts", "accounts"] = Form(...),
    mapping: str = Form(..., description="JSON object {crm_field: file_column}"),
    on_duplicate: Literal["update", "skip"] = Form("update"),
    default_lifecycle: str = Form("lead"),
    default_source: str = Form("import"),
    default_role: str = Form("customer_contact"),
    default_account_type: str = Form("prospect"),
    brand_id: int | None = Form(None),
    owner_id: int | None = Form(None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_writer),
):
    try:
        mp: dict = json.loads(mapping)
    except json.JSONDecodeError:
        raise HTTPException(422, "mapping must be JSON") from None
    headers, rows = read_table(file)
    unknown = [c for c in mp.values() if c and c not in headers]
    if unknown:
        raise HTTPException(422, f"Mapped columns not in file: {unknown}")
    brands = {b.code.lower(): b.id for b in db.scalars(select(Brand))} | {
        b.name.lower(): b.id for b in db.scalars(select(Brand))}
    job = ImportJob(entity=entity, filename=file.filename or "upload", total=len(rows), errors=[],
                    created=0, updated=0, skipped=0,
                    created_by_id=user.id)
    errors: list[dict] = []
    owner = owner_id or user.id

    def val(row: dict, field: str) -> str:
        col = mp.get(field)
        return (row.get(col) or "").strip() if col else ""

    for idx, row in enumerate(rows, start=2):  # row 1 is the header
        try:
            with db.begin_nested():
                if entity == "contacts":
                    result = _import_contact(db, row, val, on_duplicate, default_lifecycle, default_source,
                                             default_role, default_account_type, brand_id, brands, owner)
                else:
                    result = _import_account(db, row, val, on_duplicate, default_account_type, owner)
            setattr(job, result, getattr(job, result) + 1)
        except ValueError as exc:
            job.skipped += 1
            errors.append({"row": idx, "error": str(exc)})
    job.errors = errors[:500]
    db.add(job)
    db.commit()
    return job


def _import_contact(db, row, val, on_duplicate, d_lifecycle, d_source, d_role, d_acc_type, brand_id, brands,
                    owner) -> str:
    email = val(row, "email").lower() or None
    if email:
        try:
            _email.validate_python(email)
        except ValidationError:
            raise ValueError(f"Invalid email '{email}'") from None
    first, last = val(row, "first_name"), val(row, "last_name")
    full = val(row, "full_name")
    if full and not (first or last):
        parts = full.split(None, 1)
        first, last = parts[0], parts[1] if len(parts) > 1 else ""
    if not (email or first or last):
        raise ValueError("Row has neither a name nor an email")

    existing = find_contact_by_email(db, email)
    if existing and on_duplicate == "skip":
        return "skipped"

    company = val(row, "company")
    domain = _domain_from(val(row, "website")) or email_domain(email)
    account = find_or_create_account(db, company or None, domain if company or domain else None,
                                     account_type=d_acc_type, owner_id=owner,
                                     country=val(row, "country") or None) if (company or domain) else None
    fields = {
        "first_name": first or None, "last_name": last or None, "email": email,
        "phone": val(row, "phone") or None, "mobile": val(row, "mobile") or None,
        "job_title": val(row, "job_title") or None, "country": val(row, "country") or None,
        "linkedin": val(row, "linkedin") or None, "notes": val(row, "notes") or None,
    }
    lifecycle = _enum(val(row, "lifecycle"), Lifecycle, d_lifecycle) if val(row, "lifecycle") else d_lifecycle
    source = _enum(val(row, "source"), LeadSource, d_source) if val(row, "source") else d_source
    role = _enum(val(row, "contact_role"), ContactRole, d_role) if val(row, "contact_role") else d_role
    b_id = brands.get(val(row, "brand").lower(), brand_id) if val(row, "brand") else brand_id
    consent = _truthy(val(row, "consent_marketing")) if val(row, "consent_marketing") else None

    if existing:
        for k, v in fields.items():  # fill blanks / overwrite with non-empty values
            if v:
                setattr(existing, k, v)
        if account and not existing.account_id:
            existing.account_id = account.id
        if val(row, "lifecycle"):
            existing.lifecycle = lifecycle
        if consent is not None:
            existing.consent_marketing = consent
        return "updated"
    c = Contact(**fields, account_id=account.id if account else None, lifecycle=lifecycle, source=source,
                role=role, brand_id=b_id, owner_id=owner, consent_marketing=bool(consent), extra={})
    db.add(c)
    db.flush()
    return "created"


def _import_account(db, row, val, on_duplicate, d_acc_type, owner) -> str:
    name = val(row, "name")
    if not name:
        raise ValueError("Account name is empty")
    domain = _domain_from(val(row, "website"))
    existing = None
    if domain:
        existing = db.scalar(select(Account).where(Account.domain == domain))
    if not existing:
        existing = db.scalar(select(Account).where(func.lower(Account.name) == name.lower()))
    if existing and on_duplicate == "skip":
        return "skipped"
    fields = {k: (val(row, k) or None) for k in ("industry", "country", "region", "city", "address", "phone",
                                                 "vat_number", "notes")}
    fields["website"] = val(row, "website") or None
    acc_type = _enum(val(row, "account_type"), AccountType, d_acc_type) if val(row, "account_type") else d_acc_type
    if existing:
        for k, v in fields.items():
            if v:
                setattr(existing, k, v)
        if domain and not existing.domain:
            existing.domain = domain
        if val(row, "account_type"):
            existing.account_type = acc_type
        return "updated"
    db.add(Account(name=name, domain=domain, account_type=acc_type, owner_id=owner, extra={}, **fields))
    db.flush()
    return "created"


@router.get("/imports", response_model=list[ImportResult])
def import_history(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return db.scalars(select(ImportJob).order_by(ImportJob.id.desc()).limit(50)).all()


# ---------------- Export ----------------
def _respond(headers: list[str], rows: list[list], filename: str, fmt: str):
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                max(len(str(c.value or "")) for c in col) + 2, 50)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        s = io.StringIO()
        w = csv.writer(s)
        w.writerow(headers)
        w.writerows(rows)
        buf = io.BytesIO(("﻿" + s.getvalue()).encode("utf-8"))
        media = "text/csv"
    return StreamingResponse(buf, media_type=media,
                             headers={"Content-Disposition": f'attachment; filename="{filename}.{fmt}"'})


@router.get("/export/contacts")
def export_contacts(fmt: Literal["csv", "xlsx"] = "xlsx", q: str | None = None,
                    lifecycle: list[str] | None = Query(None), source: list[str] | None = Query(None),
                    account_id: int | None = None, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    stmt = contacts_router.build_query(q, account_id, lifecycle, None, source).order_by(Contact.id)
    headers = ["ID", "First name", "Last name", "Email", "Phone", "Mobile", "Job title", "Account", "Role",
               "Lifecycle", "Source", "Country", "LinkedIn", "Marketing consent", "Unsubscribed", "Owner", "Created"]
    rows = [[c.id, c.first_name, c.last_name, c.email, c.phone, c.mobile, c.job_title,
             c.account.name if c.account else None, c.role, c.lifecycle, c.source, c.country, c.linkedin,
             "yes" if c.consent_marketing else "no", "yes" if c.unsubscribed else "no",
             c.owner.full_name if c.owner else None, c.created_at.strftime("%Y-%m-%d")]
            for c in db.scalars(stmt).unique()]
    return _respond(headers, rows, "contacts", fmt)


@router.get("/export/accounts")
def export_accounts(fmt: Literal["csv", "xlsx"] = "xlsx", q: str | None = None,
                    account_type: list[str] | None = Query(None), db: Session = Depends(get_db),
                    _: User = Depends(get_current_user)):
    stmt = accounts_router.build_query(q, account_type).order_by(Account.name)
    headers = ["ID", "Name", "Type", "Domain", "Website", "Industry", "Country", "Region", "City", "Phone",
               "VAT", "Owner", "Created"]
    rows = [[a.id, a.name, a.account_type, a.domain, a.website, a.industry, a.country, a.region, a.city, a.phone,
             a.vat_number, a.owner.full_name if a.owner else None, a.created_at.strftime("%Y-%m-%d")]
            for a in db.scalars(stmt)]
    return _respond(headers, rows, "accounts", fmt)


@router.get("/export/opportunities")
def export_opportunities(fmt: Literal["csv", "xlsx"] = "xlsx", pipeline_id: int | None = None,
                         brand_id: int | None = None, status: str | None = None, q: str | None = None,
                         db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    stmt = opps_router.build_query(q, pipeline_id, None, brand_id, None, None, None, status).order_by(
        Opportunity.id)
    headers = ["ID", "Title", "Account", "Brand", "Stage", "Probability %", "Deal type", "Lead source",
               "Referring platform", "Partner contact", "Route to market", "Currency", "TCV", "ARR", "Weighted",
               "Expected close", "Contract start", "Contract end", "Term (months)", "Owner", "Products"]
    rows = []
    for o in db.scalars(stmt).unique():
        rows.append([o.id, o.title, o.account.name, o.brand.name, o.stage.name, o.effective_probability,
                     o.deal_type, o.lead_source, o.referring_platform.name if o.referring_platform else None,
                     o.partner_contact.full_name if o.partner_contact else None, o.route_to_market, o.currency,
                     float(o.amount_tcv), float(o.amount_arr), float(o.weighted_amount), o.expected_close_date,
                     o.contract_start, o.contract_end, o.term_months, o.owner.full_name if o.owner else None,
                     "; ".join(f"{li.name} x{li.quantity.normalize()}" for li in o.line_items)])
    return _respond(headers, rows, "opportunities", fmt)


# ---------------- Freshsales migration ----------------
@router.post("/migrations/freshsales")
def migrate_freshsales(
    files: list[UploadFile] = File(..., description="Freshsales exports: contacts, accounts and/or deals"),
    commit: bool = Form(False),
    default_brand: Literal["METRICSHUB", "SENTRY"] = Form("METRICSHUB"),
    default_currency: str = Form("USD"),
    stage_map: str = Form("{}", description='JSON {"Freshsales stage": "CRM stage"}'),
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    from ..migration.freshsales import from_tables, load

    try:
        overrides = json.loads(stage_map or "{}")
    except json.JSONDecodeError:
        raise HTTPException(422, "stage_map must be JSON") from None
    tables = []
    for f in files:
        headers, rows = read_table(f)
        tables.append((f.filename or "file", headers, rows))
    ds = from_tables(tables)
    rep = load(db, ds, dry_run=not commit, default_brand=default_brand,
               default_currency=default_currency.upper()[:3], stage_overrides=overrides)
    if commit:
        db.add(ImportJob(entity="freshsales", filename=", ".join(t[0] for t in tables),
                         total=len(ds.accounts) + len(ds.contacts) + len(ds.deals),
                         created=sum(c.get("created", 0) for c in rep.as_dict()["counts"].values()),
                         updated=sum(c.get("updated", 0) for c in rep.as_dict()["counts"].values()),
                         skipped=0, errors=rep.warnings[:100], created_by_id=user.id))
        db.commit()
    return {"files": [{"name": t[0], "rows": len(t[2])} for t in tables], **rep.as_dict()}
