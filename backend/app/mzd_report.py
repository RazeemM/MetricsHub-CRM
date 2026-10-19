"""Reader for BMC MarketZone "Reseller Fee Report to Sentry Software" (quarterly), as PDF or as the Detail .xlsx.

BMC has used several layouts:
  * 2020 - Sept 2024: Revenue Type | Product Name | Invoice # | Invoice Date | Invoice Type | [Order Transaction Type]
    | Install Company | Install Country | Billing Company | Ref Inv # | Number of copies | Net | Fee | Maintenance Begin
    | Maintenance End | Term (days)
  * Dec 2024 onward: Inv/Cm Number (INV..., CM... or Oracle numbers) | Inv Cm Date | Ordering Company | Sold To Country
    | Billing Company | [Related Invoice] | Product Name | Charge (Term/SaaS/Support/License) | Service Start | Service End
    | Term (days) | Quantity | Net | Fee
Columns are recognised by their header names, so their order does not matter.

The Detail .xlsx is exact and preferred. A PDF has no table cells: columns are found per page from the vertical gaps
between the data rows, named from the header words above them, and cells that run together are split by content.
Both are cross-checked against the "... Total:" rows of the report.
"""
from __future__ import annotations

import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

DATE = re.compile(r"\d{1,2}/\d{1,2}/\d{2,4}")
MONEY = re.compile(r"\(?-?[\d,]*\d\.\d{2}\)?")
NUM_TOKEN = re.compile(r"^\(?-?[\d,]*\d(?:\.\d+)?\)?$")
TOTAL = re.compile(r"(?<![A-Za-z])total\b[^:]{0,60}:\s*([\d,.()\- ]+)$", re.I)
INT_OK = re.compile(r"^\(?-?\d{1,3}(?:,\d{3})*(?:\.\d+)?\)?$|^\(?-?\d+(?:\.\d+)?\)?$")
DEC_OK = re.compile(r"^\(?-?\d{1,3}(?:,\d{3})*\.\d{2}\)?$|^\(?-?\d+\.\d{2}\)?$")
CHARGE_PREFIX = re.compile(r"^([A-Z]{4,}(?:\s*Rev\b)?)\s*(.*)$")
DOC_TYPES = ("credit memo", "credit", "invoice", "debit memo")
GLUED = re.compile(r"^(?:(?:Invoice|Credit Memo|Debit Memo)\s|(?:ADD|REN|RENEWAL|INITIAL|PERP|UPGRADE)\s|(?:SUPPORT|RENEWAL SUPPORT|INITIAL SUPPORT)[A-Z])")
MAX_TEXT = {"customer": 255, "billing": 255, "product": 255, "country": 80, "doc": 80}

# formatted fields can be split out of merged cells by their shape
FORMATTED = {"doc_date": "date", "start": "date", "end": "date", "qty": "num", "net": "num", "fee": "num", "days": "num"}


class ReportError(ValueError):
    pass


@dataclass
class ReportLine:
    doc_number: str
    doc_date: dt.date | None
    customer: str
    country: str
    billing_company: str
    product: str
    license_type: str
    start: dt.date | None
    end: dt.date | None
    term_days: int | None
    quantity: Decimal
    net: Decimal
    fee: Decimal
    page: int

    @property
    def key(self) -> str:
        return "|".join(str(x) for x in ("bmc_mzd", self.doc_number, self.product.lower(), self.start, self.end,
                                         self.quantity, self.net))


@dataclass
class Report:
    title: str
    lines: list[ReportLine]
    warnings: list[str]  # problems: lines skipped or totals that do not add up. A report with any is not imported.
    checked_totals: int = 0
    source: str = "pdf"

    @property
    def reliable(self) -> bool:
        return not self.warnings


# ---------------------------------------------------------------- values
def field_for(header: str) -> str | None:
    """Map a column header (any layout) to a field name."""
    h = re.sub(r"\s+", " ", str(header or "").lower()).strip()
    if not h:
        return None
    if "product name" in h:
        return "product"
    if "ordering company" in h or "install company" in h or "subscription owner" in h:
        return "customer"
    if "billing company" in h or "invoice owner" in h:
        return "billing"
    if "country" in h:
        return "country"
    if "quantity" in h or "copies" in h:
        return "qty"
    if "ref inv" in h or "related" in h:
        return "overflow"  # usually empty: long company names run into it
    if "revenue type" in h or "charge" in h:
        return "charge"
    if "start" in h or "begin" in h:
        return "start"
    if "term" in h and "day" in h:
        return "days"
    if "end date" in h or "end" in h.split():
        return "end"
    if "net" in h.split() or "net amount" in h:
        return "net"
    if h == "fee":
        return "fee"
    if "transaction type" in h:
        return "order_type"
    if "invoice type" in h or h == "type":
        return "doc_type"
    if "number" in h or "invoice #" in h or h == "invoice":
        return "doc"
    if "date" in h:
        return "doc_date"
    return None


def _date(v) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = DATE.search(str(v))
    if not m:
        return None
    for fmt in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return dt.datetime.strptime(m.group(0), fmt).date()
        except ValueError:
            pass
    return None


def _num(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    try:
        d = Decimal(s.strip("()"))
    except InvalidOperation:
        return None
    return -d if neg else d


def _license_type(charge: str) -> str:
    c = (charge or "").lower()
    if "saas" in c:
        return "saas"
    if "support" in c or "maint" in c:
        return "support"
    if "perpetual" in c or "license" in c or "licence" in c:
        return "perpetual"
    if "subscr" in c or "term" in c:
        return "term"
    return "other"


def _clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def make_line(f: dict, page: int) -> ReportLine:
    """Build a line from {field: raw value}; raises ReportError when the row is not usable."""
    product, charge = _clean(f.get("product")), _clean(f.get("charge"))
    if not charge and (m := CHARGE_PREFIX.match(product)):  # "PERPETUAL Rev Sentry ..." in one cell
        charge, product = m.group(1), m.group(2)
    if (m := CHARGE_PREFIX.match(product)) and _license_type(m.group(1)) != "other":
        product = m.group(2)
    product = re.sub(r"^Rev\s+", "", product)  # old layout: BMC's revenue code "Rev" in the product column
    doc = _clean(f.get("doc"))
    qty, net, fee = _num(f.get("qty")), _num(f.get("net")), _num(f.get("fee"))
    if not product or not doc:
        raise ReportError("no product or invoice number")
    if qty is None or net is None:
        raise ReportError("no quantity or amount")
    if "credit" in _clean(f.get("doc_type")).lower() and qty > 0 and net < 0:
        qty = -qty
    billing = re.sub(r"\s+[A-Z]{0,4}\d{6,}$", "", _clean(f.get("billing")))
    line = ReportLine(
        doc_number=doc, doc_date=_date(f.get("doc_date")), customer=_clean(f.get("customer")),
        country=_clean(f.get("country")), billing_company=billing, product=product,
        license_type=_license_type(charge), start=_date(f.get("start")), end=_date(f.get("end")),
        term_days=int(d) if (d := _num(f.get("days"))) is not None else None, quantity=qty,
        net=net.quantize(Decimal("0.01")), fee=(fee or Decimal(0)).quantize(Decimal("0.01")), page=page)
    if not line.customer:
        raise ReportError("no customer")
    for value in (line.customer, line.billing_company):
        if GLUED.match(value):
            raise ReportError(f"company name '{value[:40]}' starts with text from another column")
    if line.end is None and line.license_type != "perpetual":
        raise ReportError("no service end date")
    for name, limit in MAX_TEXT.items():
        value = {"customer": line.customer, "billing": line.billing_company, "product": line.product,
                 "country": line.country, "doc": line.doc_number}[name]
        if len(value) > limit:
            raise ReportError(f"{name} is {len(value)} characters long: the row was not read correctly")
    return line


def total_numbers(text: str) -> tuple[Decimal, Decimal] | None:
    """Quantity and net amount of a '... Total: qty net fee' row. Some PDFs space digits apart ('1 52,568.79'):
    the one way of joining the fragments into an integer and two amounts is used."""
    m = TOTAL.search(text)
    if not m:
        return None
    parts = m.group(1).split()
    n = len(parts)
    if n < 3 or n > 9:
        return None
    from itertools import combinations
    for cuts in combinations(range(1, n), 2):
        a, b = cuts
        q, net, fee = "".join(parts[:a]), "".join(parts[a:b]), "".join(parts[b:])
        if INT_OK.match(q) and DEC_OK.match(net) and DEC_OK.match(fee):
            return _num(q), _num(net)
    return None


class _Totals:
    """Checks each '... Total: qty net fee' row against the lines read since the previous one."""

    def __init__(self, warnings: list[str]):
        self.warnings, self.start, self.checked = warnings, 0, 0

    def check(self, lines: list[ReportLine], qty: Decimal, net: Decimal, where: str, label: str):
        section = lines[self.start:]
        self.start = len(lines)
        if not section:
            return
        got_q, got_n = sum(x.quantity for x in section), sum(x.net for x in section)
        # some sections have no total row of their own: then the total covers only the last lines
        q, n, ok = Decimal(0), Decimal(0), False
        for k, x in enumerate(reversed(section), 1):
            q, n = q + x.quantity, n + x.net
            if q == qty and abs(n - net) <= Decimal("0.01") * k + Decimal("0.02"):
                ok = True
                break
        if ok:
            self.checked += 1
        else:
            self.warnings.append(f"{where}: '{label}' is {qty} / {net:,.2f} but the lines read add up to "
                                 f"{got_q} / {got_n:,.2f}")


# ---------------------------------------------------------------- xlsx
def parse_xlsx(data: bytes) -> Report:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ReportError(f"Could not read the Excel file: {exc}") from None
    lines: list[ReportLine] = []
    warnings: list[str] = []
    totals = _Totals(warnings)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        cols = None
        for rno, row in enumerate(rows, 1):
            if cols is None:
                mapped = [field_for(c) if isinstance(c, str) else None for c in row]
                if "product" in mapped and ("customer" in mapped or "billing" in mapped):
                    cols = mapped
                continue
            cells = [c for c in row if c not in (None, "")]
            if not cells:
                continue
            f = {}
            for name, value in zip(cols, row):
                if name and value not in (None, "") and name not in f:
                    f[name] = value
            label = None
            if not f.get("product") and not f.get("doc"):
                label = next((c for c in cells if isinstance(c, str)), None)
            if label:
                nums = [_num(c) for c in row[row.index(label) + 1:] if c not in (None, "")]
                if len(nums) >= 2 and None not in nums[:2]:
                    totals.check(lines, nums[0], nums[1], f"Row {rno}", label.strip())
                continue
            try:
                lines.append(make_line(f, rno))
            except ReportError as exc:
                warnings.append(f"Row {rno}: skipped ({exc})")
    if not lines:
        raise ReportError("No invoice lines found: is this the MarketZone Detail report?")
    return Report(title="MarketZone Detail report", lines=lines, warnings=warnings,
                  checked_totals=totals.checked, source="xlsx")


# ---------------------------------------------------------------- pdf
def _lines(words: list[dict]) -> list[list[dict]]:
    """Words grouped into printed lines by vertical overlap (totals rows mix fonts with different baselines)."""
    rows: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if rows and w["top"] < max(x["bottom"] for x in rows[-1]) - 2:
            rows[-1].append(w)
        else:
            rows.append([w])
    return [sorted(r, key=lambda w: w["x0"]) for r in rows]


def _text(line: list[dict]) -> str:
    return " ".join(w["text"] for w in line)


def _is_data(text: str) -> bool:
    return bool(DATE.search(text) and MONEY.search(text)) and total_numbers(text) is None


def _columns(data: list[list[dict]], width: float, tol: int, gap: float) -> list[tuple[float, float]]:
    cov = [0] * (int(width) + 2)
    for line in data:
        seen = set()
        for w in line:
            seen.update(range(int(w["x0"]), min(int(w["x1"]) + 1, len(cov))))
        for x in seen:
            cov[x] += 1
    cols, start, run, last = [], None, 0, 0
    for x, c in enumerate(cov):
        if c > tol:
            start = x if start is None else start
            last, run = x, 0
        elif start is not None:
            run += 1
            if run >= gap:
                cols.append((start, last + 1))
                start, run = None, 0
    if start is not None:
        cols.append((start, last + 1))
    return cols


def _phrases(words: list[dict]) -> list[dict]:
    """Header words grouped into column titles: words close together on a line, or stacked right above each other
    ("Inv Service" / "End Date"), form one title. Returns [{text, x0, x1}] left to right."""
    if not words:
        return []
    h = sorted(w["bottom"] - w["top"] for w in words)[len(words) // 2]
    parent = list(range(len(words)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i, a in enumerate(words):
        for j in range(i + 1, len(words)):
            b = words[j]
            same_line = abs(a["top"] - b["top"]) < 0.5 * h
            if same_line:
                gap = max(a["x0"], b["x0"]) - min(a["x1"], b["x1"])
                joined = gap <= 0.45 * h
            else:
                near = abs(a["top"] - b["top"]) <= 1.7 * h
                ca, cb = (a["x0"] + a["x1"]) / 2, (b["x0"] + b["x1"]) / 2
                joined = near and (a["x0"] - 1 <= cb <= a["x1"] + 1 or b["x0"] - 1 <= ca <= b["x1"] + 1)
            if joined:
                parent[find(i)] = find(j)
    groups: dict[int, list[dict]] = {}
    for i, w in enumerate(words):
        groups.setdefault(find(i), []).append(w)
    out = []
    for ws in groups.values():
        ws.sort(key=lambda w: (round(w["top"] / (0.5 * h)), w["x0"]))
        out.append({"text": " ".join(w["text"] for w in ws), "x0": min(w["x0"] for w in ws),
                    "x1": max(w["x1"] for w in ws)})
    return sorted(out, key=lambda p: p["x0"])


def _column_fields(cols, phrases) -> list[list[str]]:
    """Fields of each column: each header title goes to the column it overlaps most (or the nearest one)."""
    out: list[list[str]] = [[] for _ in cols]
    for p in phrases:
        f = field_for(p["text"])
        if not f:
            continue
        cx = (p["x0"] + p["x1"]) / 2
        ov = [max(0.0, min(p["x1"], b) - max(p["x0"], a)) for a, b in cols]
        k = max(range(len(cols)), key=lambda i: (ov[i], -abs((cols[i][0] + cols[i][1]) / 2 - cx)))
        if f not in out[k]:
            out[k].append(f)
    return out


def _score(cols, phrases) -> float:
    fields = _column_fields(cols, phrases)
    found = {f for fs in fields for f in fs if f != "overflow"}
    return len(found) - 0.3 * sum(1 for fs in fields if len([f for f in fs if f != "overflow"]) > 1)


def _cell_words(line: list[dict], cols) -> list[list[dict]]:
    """Each line's words per column; a word that crosses a column edge is cut there (using its characters)."""
    cells: list[list[dict]] = [[] for _ in cols]

    def col_of(x):
        for i, (a, b) in enumerate(cols):
            if a - 0.5 <= x <= b + 0.5:
                return i
        return min(range(len(cols)), key=lambda i: min(abs(x - cols[i][0]), abs(x - cols[i][1])))

    for w in line:
        parts: dict[int, list[dict]] = {}
        for c in w.get("chars") or []:
            parts.setdefault(col_of((c["x0"] + c["x1"]) / 2), []).append(c)
        if len(parts) <= 1:
            cells[col_of((w["x0"] + w["x1"]) / 2)].append(w)
            continue
        for k, cs in parts.items():
            cells[k].append({"text": "".join(c["text"] for c in cs), "x0": cs[0]["x0"], "x1": cs[-1]["x1"],
                             "top": w["top"], "bottom": w["bottom"]})
    return cells


CHARGE_CODE = re.compile(r"^(ONPREMSUBSCRIPT|PERPETUAL|SUBSCRIPTION|MAINTENANCE|LICENSE|SUPPORT|SAAS|TERM|"
                         r"License|Support|SaaS|Term)(\s*Rev\b)?\s*", re.I)


def _valid_date(t: str) -> bool:
    return DATE.fullmatch(t) is not None and _date(t) is not None


def _take_date(tokens: list[str]) -> str | None:
    """Remove and return the first date; also cuts one off the end of a token it is glued to ('23871693/20/2020')."""
    for i, t in enumerate(tokens):
        if _valid_date(t):
            return tokens.pop(i)
        for k in range(1, len(t) - 5):
            if t[:k].isdigit() and _valid_date(t[k:]):
                tokens[i] = t[:k]
                return t[k:]
    return None


def _split_cell(words: list[dict], fields: list[str]) -> dict:
    """Split one cell that holds several fields (columns that ran together on this page), by content."""
    text = " ".join(w["text"] for w in words)
    if len(fields) == 1:
        return {fields[0]: text}
    out: dict = {}
    fields = list(fields)
    if "charge" in fields and "product" in fields and (m := CHARGE_CODE.match(text)):
        out["charge"] = m.group(0).strip()
        text = text[m.end():]
        fields.remove("charge")
    if "doc_type" in fields:
        low = text.lower()
        if dtype := next((d for d in DOC_TYPES if low.startswith(d)), None):
            out["doc_type"], text = text[:len(dtype)], text[len(dtype):].strip()
        fields.remove("doc_type")
    tokens = text.split()
    # dates and numbers are taken by their shape
    for f in [f for f in fields if f in FORMATTED]:
        pick = _take_date(tokens) if FORMATTED[f] == "date" else next(
            (t for t in tokens if NUM_TOKEN.match(t)), None)
        if pick is not None:
            out[f] = pick
            if FORMATTED[f] != "date":
                tokens.remove(pick)
        fields.remove(f)
    if "doc" in fields and tokens and re.fullmatch(r"[A-Z]*\d{5,}", tokens[0]):
        out["doc"] = tokens.pop(0)
        fields.remove("doc")
    if "order_type" in fields:
        if tokens and tokens[0].isupper() and len(tokens[0]) <= 12:
            out["order_type"] = tokens.pop(0)
        fields.remove("order_type")
    fields = [f for f in fields if f != "overflow"]
    if len(fields) == 1:
        out[fields[0]] = " ".join(tokens)
    elif len(fields) > 1 and tokens:
        # company / country / company: cut at the widest gaps between the remaining words
        rest = [w for w in words if w["text"] in tokens]
        if len(rest) != len(tokens):
            rest = [{"text": t, "x0": i * 10.0, "x1": i * 10.0 + 9} for i, t in enumerate(tokens)]
        gaps = sorted(range(1, len(rest)), key=lambda i: rest[i]["x0"] - rest[i - 1]["x1"], reverse=True)
        bounds = [0, *sorted(gaps[:len(fields) - 1]), len(rest)]
        for f, (a, b) in zip(fields, zip(bounds, bounds[1:])):
            out[f] = " ".join(w["text"] for w in rest[a:b])
    return out


def parse_pdf(data: bytes) -> Report:
    import pdfplumber  # imported lazily: only needed for report imports

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - pdfminer raises many types
        raise ReportError(f"Could not read the PDF: {exc}") from None
    lines: list[ReportLine] = []
    warnings: list[str] = []
    totals = _Totals(warnings)
    title = ""
    with pdf:
        first_text = (pdf.pages[0].extract_text() or "") if pdf.pages else ""
        if "Reseller Fee Report" not in first_text:
            raise ReportError("This does not look like a BMC reseller fee report")
        if "Marketzone" not in first_text and "MarketZone" not in first_text:
            raise ReportError("This is not a MarketZone report (OEM reports are not supported)")
        title = first_text.splitlines()[0].strip()
        for pno, page in enumerate(pdf.pages, 1):
            plines = _lines(page.extract_words(return_chars=True))
            anchor = next((i for i, ln in enumerate(plines) if "Product" in _text(ln) and "Name" in _text(ln)), None)
            if anchor is None:
                continue
            data_idx = [i for i in range(anchor + 1, len(plines)) if _is_data(_text(plines[i]))]
            if not data_idx:
                continue
            top = plines[anchor][0]["top"] - 40
            header_words = [w for ln in plines[max(0, anchor - 6):data_idx[0]] for w in ln if w["top"] >= top]
            data = [plines[i] for i in data_idx]
            heights = sorted(w["bottom"] - w["top"] for ln in data for w in ln)
            h = heights[len(heights) // 2]
            candidates = [_columns(data, page.width, tol, max(1.5, g * h))
                          for tol in range(0, 1 + len(data) // 15) for g in (0.45, 0.8, 1.2)]
            phrases = _phrases(header_words)
            cols = max(candidates, key=lambda c: _score(c, phrases))
            fields = _column_fields(cols, phrases)
            flat = [f for fs in fields for f in fs]
            if "end" not in flat and "start" in flat:
                seen_start = False
                for fs in fields:
                    if "start" in fs:
                        seen_start = True
                    elif seen_start and fs == ["doc_date"]:
                        fs[0] = "end"
                        break
            if "product" not in {f for fs in fields for f in fs}:
                warnings.append(f"Page {pno}: could not find the columns")
                continue
            prev_line: ReportLine | None = None
            prev_bottom = None
            data_set = set(data_idx)
            for i in range(anchor + 1, len(plines)):
                ln = plines[i]
                text = _text(ln)
                if (tn := total_numbers(text)) is not None:
                    totals.check(lines, tn[0], tn[1], f"Page {pno}", text[:60])
                    prev_line = None
                    continue
                if i not in data_set:
                    # a wrapped cell continues just below its row, inside a text column
                    if prev_line and prev_bottom is not None and ln[0]["top"] - prev_bottom < 1.5 * h \
                            and len(text) < 80 and not DATE.search(text):
                        attr = {"customer": "customer", "billing": "billing_company", "product": "product"}
                        left_text = None
                        for fs, ws in zip(fields, _cell_words(ln, cols)):
                            target = next((x for x in fs if x in attr), None)
                            if target is None and fs == ["overflow"]:
                                target = left_text
                            if ws and target:
                                setattr(prev_line, attr[target], getattr(prev_line, attr[target]) + " "
                                        + " ".join(w["text"] for w in ws))
                            if fs and fs[-1] in attr:
                                left_text = fs[-1]
                        prev_bottom = max(w["bottom"] for w in ln)
                    else:
                        prev_line = None
                    continue
                f: dict = {}
                last_text, last_x1 = None, None
                for fs, ws in zip(fields, _cell_words(ln, cols)):
                    if fs == ["overflow"]:
                        # usually empty: a long company name that runs on into it belongs to the column on its left
                        extra = " ".join(w["text"] for w in ws)
                        if extra and last_text in f and not re.fullmatch(r"[A-Z]*\d{5,}", extra):
                            glued = last_x1 is not None and ws[0]["x0"] - last_x1 < 0.5
                            f[last_text] = f[last_text] + ("" if glued else " ") + extra
                        continue
                    if ws and fs:
                        for k, v in _split_cell(ws, [x for x in fs if x != "overflow"]).items():
                            f.setdefault(k, v)
                        last_text = fs[-1] if fs[-1] in ("customer", "billing", "country", "product") else None
                        last_x1 = ws[-1]["x1"]
                try:
                    prev_line = make_line(f, pno)
                    lines.append(prev_line)
                    prev_bottom = max(w["bottom"] for w in ln)
                except ReportError as exc:
                    warnings.append(f"Page {pno}: skipped '{text[:90]}' ({exc})")
                    prev_line = None
    if not lines:
        raise ReportError("No invoice lines found in this report")
    return Report(title=title, lines=lines, warnings=warnings, checked_totals=totals.checked, source="pdf")


def parse(data: bytes, filename: str = "") -> Report:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return parse_xlsx(data)
    if data[:5] == b"%PDF-" or name.endswith(".pdf"):
        return parse_pdf(data)
    raise ReportError("Upload the report as PDF or as the Detail .xlsx")
