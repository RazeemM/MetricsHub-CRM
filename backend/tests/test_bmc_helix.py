"""BMC → BMC Helix address reconciliation workbook: move, merge, create, and nothing more on a second run."""
import io

from openpyxl import Workbook

from app.db import SessionLocal
from app.models import Activity, Contact, ListMember

HEADER = ["Ligne Excel", "ID original", "Prénom original", "Nom original", "E-mail original", "Nom dans le mail",
          "E-mail Helix observé", "Compte original", "Statut", "Candidats", "Preuve", "Méthode", "Date source UTC"]
HELIX_HEADER = ["Ligne Excel", "ID original", "Prénom original", "Nom original", "E-mail original",
                "Compte original", "Statut"]


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append([])
    ws.append([f"{len(rows)} lignes"])  # the workbook's title rows above the header
    ws.append(["Note"])
    ws.append([])
    ws.append(header)
    for r in rows:
        ws.append(r)
    if not rows:
        ws.append(["Aucune ligne dans cette catégorie."])


def _workbook() -> bytes:
    wb = Workbook()
    wb.active.title = "Résumé"
    m = lambda i, first, last, old, new, acc: [i, str(9000 + i), first, last, old, f"{last}, {first}", new, acc,  # noqa: E731
                                               "PROBABLE_MAIL_MAI_2026", 1, "Probable, sur nom complet uniquement",
                                               "Nom complet normalisé", "2026-05-04"]
    _sheet(wb, "Correspondances", HEADER, [
        m(1, "Mona", "Moved", "mona_moved@bmc.com", "mona@helixops.ai", "BMC Helix"),     # old only → moved
        m(2, "Dina", "Double", "dina_double@bmc.com", "dina@helixops.ai", "BMC Software"),  # both → merged
        m(3, "Nora", "New", "nora_new@bmc.com", "nora@helixops.ai", "Unknown Corp Ltd"),  # neither → created
        m(4, "Hugo", "Helix", "hugo_helix@bmc.com", "hugo@helixops.ai", "BMC"),           # new only → noted
    ])
    _sheet(wb, "Non trouvés", HEADER, [
        [5, "9005", "Fay", "Found", "fay_found@bmc.com", None, None, "BMC", "NON_TROUVE_MAIL_MAI_2026", 0],
        [6, "9006", "Gil", "Gone", "gil_gone@bmc.com", None, None, "BMC Software", "NON_TROUVE_MAIL_MAI_2026", 0],
    ])
    _sheet(wb, "Ambigus", HEADER, [
        [7, "9007", "Amy", "Ambiguous", "amy_amb@bmc.com", None, "amy@helixops.ai", "BMC", "AMBIGU", 2],
    ])
    _sheet(wb, "Déjà Helix", HELIX_HEADER, [[8, "9008", "Deb", "Deja", "deb@helixops.ai", "BMC Helix",
                                            "DEJA_HELIX_DANS_FICHIER"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _run(client, commit, **form):
    return client.post("/api/migrations/bmc-helix", files={"file": ("rapprochement.xlsx", _workbook())},
                       data={"commit": str(commit).lower(), **form})


def _by_email(email):
    with SessionLocal() as db:
        c = db.query(Contact).filter(Contact.email == email).one_or_none()
        if c:
            db.expunge(c)
        return c


def test_bmc_helix_reconciliation(client):
    post = lambda **kw: client.post("/api/contacts", json=kw).json()  # noqa: E731
    mona = post(first_name="Mona", last_name="Moved", email="mona_moved@bmc.com")
    dina_old = post(first_name="Dina", email="dina_double@bmc.com", phone="+1 555 0100")
    dina_new = post(first_name="Dina", last_name="Double", email="dina@helixops.ai")
    post(first_name="Hugo", last_name="Helix", email="hugo@helixops.ai")
    post(first_name="Fay", email="fay_found@bmc.com")
    lst = client.post("/api/lists", json={"name": "BMC partner news"}).json()
    client.post(f"/api/lists/{lst['id']}/members", json={"contact_ids": [dina_old["id"]]})
    client.post("/api/activities", json={"type": "note", "subject": "Met Dina at the summit", "contact_id": dina_old["id"]})

    dry = _run(client, False).json()
    assert dry["dry_run"] and dry["rows"] == 8
    assert dry["rows_by_sheet"] == {"match": 4, "not_found": 2, "ambiguous": 1, "already_helix": 1}
    assert dry["counts"] == {"email_updated": 1, "merged": 1, "created": 3, "updated": 2, "skipped_ambiguous": 1}
    assert _by_email("mona_moved@bmc.com") and not _by_email("nora@helixops.ai")  # nothing written

    done = _run(client, True).json()
    assert done["counts"] == dry["counts"] and not done["dry_run"]

    moved = _by_email("mona@helixops.ai")
    assert moved.id == mona["id"] and moved.extra["previous_emails"] == ["mona_moved@bmc.com"]
    assert moved.extra["bmc_helix"]["status"] == "PROBABLE_MAIL_MAI_2026"
    assert moved.account_id is not None and moved.extra["bmc_helix"]["source_id"] == "9001"  # BMC platform account

    dina = _by_email("dina@helixops.ai")
    assert dina.id == dina_new["id"] and dina.phone == "+1 555 0100" and _by_email("dina_double@bmc.com") is None
    assert "dina_double@bmc.com" in dina.extra["previous_emails"]
    with SessionLocal() as db:
        assert db.query(ListMember).filter_by(list_id=lst["id"], contact_id=dina.id).count() == 1
        subjects = {a.subject for a in db.query(Activity).filter_by(contact_id=dina.id)}
        assert "Met Dina at the summit" in subjects and any(s.startswith("Merged with") for s in subjects)

    nora = _by_email("nora@helixops.ai")
    assert nora.account_id is None and nora.extra["company"] == "Unknown Corp Ltd"  # unknown accounts aren't created
    assert _by_email("gil_gone@bmc.com").role == "platform_account_manager"  # not found: kept the BMC address
    assert _by_email("deb@helixops.ai") is not None
    assert _by_email("hugo@helixops.ai").extra["previous_emails"] == ["hugo_helix@bmc.com"]
    assert _by_email("amy@helixops.ai") is None and _by_email("amy_amb@bmc.com") is None
    assert _by_email("fay_found@bmc.com").last_name == "Found"  # blanks filled, the address kept

    found = client.get("/api/contacts", params={"q": "mona_moved@bmc"}).json()["items"]  # by the old address
    assert [c["id"] for c in found] == [mona["id"]] and found[0]["previous_emails"] == ["mona_moved@bmc.com"]
    assert client.get("/api/search", params={"q": "mona_moved"}).json()["contacts"][0]["id"] == mona["id"]

    again = _run(client, True).json()  # the same workbook again: nothing left to do
    assert again["counts"] == {"unchanged": 7, "skipped_ambiguous": 1}


def test_bmc_helix_options_and_bad_files(client):
    r = _run(client, False, apply_matches="false", create_missing="false").json()
    assert "email_updated" not in r["counts"] and "created" not in r["counts"]
    bad = client.post("/api/migrations/bmc-helix", files={"file": ("x.xlsx", b"not a workbook")}, data={})
    assert bad.status_code == 422
    wb = Workbook()
    wb.active.append(["Name", "Email"])
    out = io.BytesIO()
    wb.save(out)
    other = client.post("/api/migrations/bmc-helix", files={"file": ("x.xlsx", out.getvalue())}, data={})
    assert other.status_code == 422 and "reconciliation" in other.json()["detail"]
